import argparse
import csv
import sys
import time
from pathlib import Path

import cadquery as cq
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from plate import plate

NUMERIC_FIELDS = ["length_mm", "width_mm", "height_mm", "hole_dia_mm", "fillet_r_mm"]
REQUIRED_COLUMNS = ["variant_id"] + NUMERIC_FIELDS
ALLOWED_FORMATS = {"step", "stl"}


def build_manifest_fields(formats):
    fields = ["variant_id"] + NUMERIC_FIELDS
    if "step" in formats:
        fields.append("step_file")
    if "stl" in formats:
        fields.append("stl_file")
    fields.append("status")
    return fields


def read_csv_rows(csv_path):
    with open(csv_path, newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        fieldnames = reader.fieldnames or []
        rows = list(enumerate(reader, start=2))
    return fieldnames, rows


def validate(fieldnames, rows):
    # Input errors abort the whole run before any geometry is built;
    # only CAD failures (in generate()) are handled per row.
    errors = []

    missing_columns = [c for c in REQUIRED_COLUMNS if c not in fieldnames]
    if missing_columns:
        for col in missing_columns:
            errors.append(f"missing column: {col}")
        return errors, None

    resolved = {}
    id_rows = {}
    auto_counter = 0

    for row_num, row in rows:
        raw_values = {field: (row.get(field) or "").strip() for field in NUMERIC_FIELDS}
        missing_fields = [field for field, raw in raw_values.items() if not raw]
        if missing_fields:
            errors.append(f"row {row_num}: missing value ({', '.join(missing_fields)})")
            continue

        bad_fields = []
        values = {}
        for field, raw in raw_values.items():
            try:
                values[field] = float(raw)
            except ValueError:
                bad_fields.append(f"{field}={raw!r}")
        if bad_fields:
            errors.append(f"row {row_num}: not a number ({', '.join(bad_fields)})")
            continue

        variant_id = (row.get("variant_id") or "").strip()
        if not variant_id:
            auto_counter += 1
            variant_id = f"V{auto_counter:03d}"

        id_rows.setdefault(variant_id, []).append(row_num)
        resolved[row_num] = {"variant_id": variant_id, **values}

    for variant_id, row_nums in id_rows.items():
        if len(row_nums) > 1:
            errors.append(f"duplicate variant_id: {variant_id} (rows: {', '.join(map(str, row_nums))})")

    return errors, resolved


def _to_number(value):
    """Returns int for whole numbers, float otherwise -- so Excel shows 60, not 60.0."""
    f = float(value)
    return int(f) if f.is_integer() else f


HEADER_FONT = Font(name="Calibri", size=11, bold=True)
HEADER_FILL = PatternFill(start_color="FFD9D9D9", end_color="FFD9D9D9", fill_type="solid")
DATA_FONT = Font(name="Calibri", size=11)
STATUS_OK_FONT = Font(name="Calibri", size=11, color="FF008000")
STATUS_FAILED_FONT = Font(name="Calibri", size=11, color="FFFF0000")
THIN_SIDE = Side(style="thin", color="FF000000")
THIN_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)


def write_manifest_xlsx(manifest_rows, xlsx_path, manifest_fields):
    wb = Workbook()
    ws = wb.active
    ws.title = "manifest"

    numeric_fields = set(NUMERIC_FIELDS)

    ws.append(manifest_fields)
    for col_idx, field in enumerate(manifest_fields, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER
        cell.alignment = Alignment(horizontal="right" if field in numeric_fields else "left")

    for row in manifest_rows:
        values = [_to_number(row[f]) if f in numeric_fields else row[f] for f in manifest_fields]
        ws.append(values)

    for row_offset, row in enumerate(manifest_rows):
        row_idx = row_offset + 2
        for col_idx, field in enumerate(manifest_fields, start=1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal="right" if field in numeric_fields else "left")
            if field == "status":
                if row["status"] == "ok":
                    cell.font = STATUS_OK_FONT
                elif row["status"].startswith("failed"):
                    cell.font = STATUS_FAILED_FONT
                else:
                    cell.font = DATA_FONT
            else:
                cell.font = DATA_FONT

    ws.freeze_panes = "A2"

    for col_idx, field in enumerate(manifest_fields, start=1):
        cell_lengths = [len(str(field))] + [len(str(row[field])) for row in manifest_rows]
        column_letter = ws.cell(row=1, column=col_idx).column_letter
        ws.column_dimensions[column_letter].width = max(cell_lengths) + 2

    wb.save(xlsx_path)


def generate(resolved, out_dir, formats):
    # A row's CAD failure is recorded as status="failed: ..." and skipped;
    # it never aborts the rest of the batch.
    step_dir = out_dir / "step"
    stl_dir = out_dir / "stl"
    preview_dir = out_dir / "preview"

    dirs_to_make = [out_dir, preview_dir]
    if "step" in formats:
        dirs_to_make.append(step_dir)
    if "stl" in formats:
        dirs_to_make.append(stl_dir)
    for d in dirs_to_make:
        d.mkdir(parents=True, exist_ok=True)

    manifest_fields = build_manifest_fields(formats)
    manifest_rows = []

    for row_num in sorted(resolved):
        data = resolved[row_num]
        variant_id = data["variant_id"]
        values = {k: v for k, v in data.items() if k != "variant_id"}

        row_result = {"variant_id": variant_id, **values}

        try:
            part = plate(
                length=values["length_mm"],
                width=values["width_mm"],
                height=values["height_mm"],
                hole_dia=values["hole_dia_mm"],
                fillet_r=values["fillet_r_mm"],
            )
            if "step" in formats:
                step_rel = f"step/{variant_id}.step"
                cq.exporters.export(part, str(out_dir / step_rel))
                row_result["step_file"] = step_rel
            if "stl" in formats:
                stl_rel = f"stl/{variant_id}.stl"
                cq.exporters.export(part, str(out_dir / stl_rel), tolerance=0.01, angularTolerance=0.05)
                row_result["stl_file"] = stl_rel
            status = "ok"
        except Exception as e:
            status = f"failed: {e}"
            if "step" in formats:
                row_result.setdefault("step_file", "")
            if "stl" in formats:
                row_result.setdefault("stl_file", "")

        row_result["status"] = status
        manifest_rows.append(row_result)

    manifest_path = out_dir / "manifest.csv"
    with open(manifest_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=manifest_fields)
        writer.writeheader()
        writer.writerows(manifest_rows)

    write_manifest_xlsx(manifest_rows, out_dir / "manifest.xlsx", manifest_fields)

    return manifest_rows, manifest_path


def main(csv_path="variants.csv", out_dir=None, formats=None):
    start_time = time.perf_counter()
    out_dir = Path(out_dir) if out_dir else Path("output")
    formats = formats or ["step", "stl"]

    fieldnames, rows = read_csv_rows(csv_path)
    errors, resolved = validate(fieldnames, rows)

    if errors:
        print(f"CSV validation failed -- {len(errors)} problem(s) found, nothing was generated:")
        for err in errors:
            print(f"  - {err}")
        sys.exit(1)

    manifest_rows, manifest_path = generate(resolved, out_dir, formats)

    elapsed = time.perf_counter() - start_time
    ok_count = sum(1 for r in manifest_rows if r["status"] == "ok")
    failed_count = len(manifest_rows) - ok_count

    print(f"Done: {len(manifest_rows)} variants processed ({ok_count} ok, {failed_count} failed).")
    print(f"Formats: {', '.join(formats)}")
    print(f"Manifest: {manifest_path}")
    print(f"Total time: {elapsed:.2f} seconds")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Generates parametric plate variants from a CSV file.")
    parser.add_argument("csv_path", nargs="?", default="variants.csv")
    parser.add_argument("out_dir", nargs="?", default="output")
    parser.add_argument("--formats", default="step,stl", help="Comma-separated: step, stl, or step,stl")
    args = parser.parse_args()

    formats_arg = [f.strip().lower() for f in args.formats.split(",") if f.strip()]
    invalid = [f for f in formats_arg if f not in ALLOWED_FORMATS]
    if invalid:
        sys.exit(f"ERROR: invalid format(s): {', '.join(invalid)} (allowed: step, stl)")
    if not formats_arg:
        sys.exit("ERROR: at least one format must be given (step, stl)")

    main(args.csv_path, args.out_dir, formats_arg)
